"""Connector-backed dataset reader implementing streaming chunk extraction."""

import csv
import json
import uuid
from collections.abc import AsyncIterator
from pathlib import Path
from typing import Any

import openpyxl

from app.connectors.domain.enums import ConnectorType
from app.connectors.domain.errors import FileProcessingError
from app.connectors.domain.models import QueryRequest
from app.connectors.infrastructure.registry import get_connector_registry
from app.ingestion.domain.errors import MaterializationError
from app.ingestion.domain.models import IngestionChunk
from app.ingestion.domain.protocols import DatasetReader


class ConnectorDatasetReader(DatasetReader):
    """Streams data chunks from registered Connectors (PostgreSQL, CSV, Excel)."""

    def __init__(
        self,
        connector_type: str,
        configuration: dict[str, Any],
        datasource_id: uuid.UUID,
        organization_id: uuid.UUID,
        target_name: str | None = None,
    ) -> None:
        self.connector_type = connector_type.lower()
        self.configuration = configuration
        self.datasource_id = datasource_id
        self.organization_id = organization_id
        self.target_name = target_name

    async def read_chunks(
        self,
        batch_size: int = 1000,
        max_rows: int = 1_000_000,
        watermark_column: str | None = None,
        last_watermark: str | None = None,
    ) -> AsyncIterator[IngestionChunk]:
        """Stream tabular data in bounded batches."""
        if self.connector_type == ConnectorType.CSV.value:
            async for chunk in self._read_csv_chunks(batch_size, max_rows):
                yield chunk
        elif self.connector_type == ConnectorType.EXCEL.value:
            async for chunk in self._read_excel_chunks(batch_size, max_rows):
                yield chunk
        elif self.connector_type == ConnectorType.POSTGRESQL.value:
            async for chunk in self._read_postgres_chunks(
                batch_size, max_rows, watermark_column, last_watermark
            ):
                yield chunk
        else:
            raise MaterializationError(f"Unsupported connector type: {self.connector_type}")

    async def _read_csv_chunks(
        self,
        batch_size: int,
        max_rows: int,
    ) -> AsyncIterator[IngestionChunk]:
        raw_path = self.configuration.get("file_path")
        if not raw_path:
            raise FileProcessingError("unknown", "Missing 'file_path' in CSV configuration")

        file_path = Path(raw_path)
        if not file_path.exists():
            raise FileProcessingError(str(file_path), f"CSV file not found: {file_path}")

        chunk_idx = 0
        current_rows: list[dict[str, Any]] = []
        total_emitted = 0

        with file_path.open(mode="r", encoding="utf-8", errors="replace") as f:
            reader = csv.DictReader(f)
            if not reader.fieldnames:
                return

            for row in reader:
                if total_emitted >= max_rows:
                    break
                current_rows.append(dict(row))
                total_emitted += 1

                if len(current_rows) >= batch_size:
                    byte_size = len(json.dumps(current_rows).encode("utf-8"))
                    yield IngestionChunk(
                        chunk_index=chunk_idx,
                        rows=current_rows,
                        row_count=len(current_rows),
                        byte_size=byte_size,
                    )
                    chunk_idx += 1
                    current_rows = []

            if current_rows:
                byte_size = len(json.dumps(current_rows).encode("utf-8"))
                yield IngestionChunk(
                    chunk_index=chunk_idx,
                    rows=current_rows,
                    row_count=len(current_rows),
                    byte_size=byte_size,
                )

    async def _read_excel_chunks(
        self,
        batch_size: int,
        max_rows: int,
    ) -> AsyncIterator[IngestionChunk]:
        raw_path = self.configuration.get("file_path")
        if not raw_path:
            raise FileProcessingError("unknown", "Missing 'file_path' in Excel configuration")

        file_path = Path(raw_path)
        if not file_path.exists():
            raise FileProcessingError(str(file_path), f"Excel file not found: {file_path}")

        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        try:
            sheet_name = self.target_name or wb.sheetnames[0]
            if sheet_name not in wb.sheetnames:
                raise FileProcessingError(
                    str(file_path), f"Sheet '{sheet_name}' not found in workbook"
                )

            ws = wb[sheet_name]
            row_iter = ws.iter_rows(values_only=True)

            # Header row
            header_row = next(row_iter, None)
            if not header_row:
                return

            headers = [
                str(col).strip() if col is not None else f"col_{i}"
                for i, col in enumerate(header_row)
            ]

            chunk_idx = 0
            current_rows: list[dict[str, Any]] = []
            total_emitted = 0

            for raw_row in row_iter:
                if total_emitted >= max_rows:
                    break
                # Construct dictionary row
                row_dict = {
                    headers[i]: raw_row[i] if i < len(raw_row) else None
                    for i in range(len(headers))
                }
                current_rows.append(row_dict)
                total_emitted += 1  # noqa: SIM113

                if len(current_rows) >= batch_size:
                    byte_size = len(json.dumps(current_rows, default=str).encode("utf-8"))
                    yield IngestionChunk(
                        chunk_index=chunk_idx,
                        rows=current_rows,
                        row_count=len(current_rows),
                        byte_size=byte_size,
                    )
                    chunk_idx += 1
                    current_rows = []

            if current_rows:
                byte_size = len(json.dumps(current_rows, default=str).encode("utf-8"))
                yield IngestionChunk(
                    chunk_index=chunk_idx,
                    rows=current_rows,
                    row_count=len(current_rows),
                    byte_size=byte_size,
                )
        finally:
            wb.close()

    async def _read_postgres_chunks(
        self,
        batch_size: int,
        max_rows: int,
        watermark_column: str | None = None,
        last_watermark: str | None = None,
    ) -> AsyncIterator[IngestionChunk]:
        connector = get_connector_registry().get(ConnectorType.POSTGRESQL.value)
        table_name = self.target_name or "data"

        offset = 0
        chunk_idx = 0
        total_emitted = 0

        while total_emitted < max_rows:
            limit = min(batch_size, max_rows - total_emitted)

            # Build query with optional watermark constraint
            if watermark_column and last_watermark is not None:
                query = (
                    f"SELECT * FROM {table_name} "
                    f"WHERE {watermark_column} > '{last_watermark}' "
                    f"ORDER BY {watermark_column} ASC LIMIT {limit} OFFSET {offset}"
                )
            else:
                query = f"SELECT * FROM {table_name} LIMIT {limit} OFFSET {offset}"

            req = QueryRequest(
                datasource_id=self.datasource_id,
                organization_id=self.organization_id,
                query=query,
                timeout_ms=30000,
                max_rows=limit,
            )

            result = await connector.execute_query(req, self.configuration)
            if not result.rows:
                break

            byte_size = len(json.dumps(result.rows, default=str).encode("utf-8"))
            yield IngestionChunk(
                chunk_index=chunk_idx,
                rows=result.rows,
                row_count=len(result.rows),
                byte_size=byte_size,
            )

            chunk_idx += 1
            total_emitted += len(result.rows)
            offset += len(result.rows)

            if len(result.rows) < limit:
                break
