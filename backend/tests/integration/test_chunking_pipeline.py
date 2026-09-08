"""Integration tests for Task 9 Document Chunking Pipeline.

Covers persistence, idempotent reprocessing, RBAC, and multi-tenancy.
"""

import io
import uuid
from collections.abc import AsyncGenerator

import docx
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.jwt import create_access_token
from app.auth.password import hash_password
from app.core.config import get_settings
from app.db.postgres import (
    create_database_engine,
    create_session_factory,
    dispose_database_engine,
)
from app.main import create_app
from app.models.document import Document, DocumentChunk
from app.models.organization import Organization
from app.models.role import OrganizationMember, Role
from app.models.user import User
from app.rbac.catalog import ROLE_ADMIN, ROLE_ANALYST, ROLE_VIEWER
from app.rbac.service import RBACService


@pytest.fixture
async def async_client() -> AsyncGenerator[AsyncClient, None]:
    """Provide AsyncClient wired to FastAPI app with PostgreSQL database state."""
    settings = get_settings()
    engine = create_database_engine(settings)
    session_factory = create_session_factory(engine)

    app = create_app()
    app.state.db_engine = engine
    app.state.db_session_factory = session_factory

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        yield client

    await dispose_database_engine(engine)


@pytest.fixture
async def db_session() -> AsyncGenerator[AsyncSession, None]:
    """Provide real PostgreSQL session for DB state setup and verification."""
    settings = get_settings()
    engine = create_database_engine(settings)
    session_factory = create_session_factory(engine)
    async with session_factory() as session:
        yield session
    await dispose_database_engine(engine)


async def _create_test_organization(session: AsyncSession, name: str | None = None) -> Organization:
    uid = uuid.uuid4().hex[:8]
    org = Organization(
        name=name or f"Chunk Org {uid}",
        slug=f"chunk-org-{uid}",
        is_active=True,
    )
    session.add(org)
    await session.flush()
    return org


async def _create_test_user(session: AsyncSession, email: str | None = None) -> User:
    uid = uuid.uuid4().hex[:8]
    if email:
        prefix, domain = email.split("@", 1)
        unique_email = f"{prefix}_{uid}@{domain}"
    else:
        unique_email = f"chunk_user_{uid}@example.com"
    user = User(
        email=unique_email,
        password_hash=hash_password("Password123!"),
        first_name="Chunk",
        last_name="Tester",
        is_active=True,
        is_verified=True,
    )
    session.add(user)
    await session.flush()
    return user


async def _ensure_rbac_seeded(session: AsyncSession) -> None:
    service = RBACService()
    await service.seed_system_rbac(session)
    await session.flush()


async def _assign_member(
    session: AsyncSession, user_id: uuid.UUID, organization_id: uuid.UUID, role_name: str
) -> OrganizationMember:
    role = (
        await session.execute(
            select(Role).where(
                Role.name == role_name,
                Role.organization_id.is_(None),
            )
        )
    ).scalar_one()

    membership = OrganizationMember(
        user_id=user_id,
        organization_id=organization_id,
        role_id=role.id,
    )
    session.add(membership)
    await session.flush()
    return membership


async def _create_test_user_with_role(
    session: AsyncSession,
    org: Organization,
    role_name: str,
    email: str | None = None,
) -> tuple[User, str]:
    await _ensure_rbac_seeded(session)
    user = await _create_test_user(session, email)
    await _assign_member(session, user_id=user.id, organization_id=org.id, role_name=role_name)
    await session.commit()
    token = create_access_token(user_id=user.id, extra_claims={"email": user.email})
    return user, token


def _create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Corporate Performance 2025", level=1)
    doc.add_paragraph("Fiscal year 2025 delivered record revenue across EMEA and APAC regions.")
    doc.add_heading("Financial Breakdown", level=2)
    doc.add_paragraph("Operating cash flow increased by 18% year-over-year.")
    doc.add_paragraph("Key Highlights:")
    doc.add_paragraph("Total revenue reached 120 million USD.", style="List Bullet")
    doc.add_paragraph("Customer count expanded to 4,500 enterprises.", style="List Bullet")

    table = doc.add_table(rows=3, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Quarter"
    hdr_cells[1].text = "Revenue"
    hdr_cells[2].text = "Margin"
    row1 = table.rows[1].cells
    row1[0].text = "Q1"
    row1[1].text = "$28M"
    row1[2].text = "24%"
    row2 = table.rows[2].cells
    row2[0].text = "Q2"
    row2[1].text = "$32M"
    row2[2].text = "26%"

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# Integration Tests
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_chunking_pipeline_end_to_end(
    async_client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    """Test Upload -> Ingest -> Chunk -> Verify PostgreSQL persistence and parent/child linkage."""
    org = await _create_test_organization(db_session)
    _, token = await _create_test_user_with_role(db_session, org, ROLE_ADMIN)
    headers = {"Authorization": f"Bearer {token}", "X-Organization-ID": str(org.id)}

    # 1. Upload DOCX document
    docx_bytes = _create_sample_docx()
    upload_res = await async_client.post(
        "/api/v1/documents",
        headers=headers,
        files={
            "file": (
                "report.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload_res.status_code == 201
    doc_id = upload_res.json()["id"]

    # 2. Ingest document synchronously
    ingest_res = await async_client.post(
        f"/api/v1/documents/{doc_id}/ingest?sync=true",
        headers=headers,
    )
    assert ingest_res.status_code == 200
    assert ingest_res.json()["status"] == "parsed"

    # 3. Chunk document synchronously
    chunk_res = await async_client.post(
        f"/api/v1/documents/{doc_id}/chunks?sync=true",
        headers=headers,
    )
    assert chunk_res.status_code == 200
    summary = chunk_res.json()
    assert summary["document_id"] == doc_id
    assert summary["organization_id"] == str(org.id)
    assert summary["total_chunks"] > 0
    assert summary["parent_chunks"] >= 1
    assert summary["child_chunks"] >= 1
    assert summary["min_tokens"] > 0
    assert summary["empty_chunks"] == 0
    assert summary["oversized_chunks"] == 0

    # 4. Check document status in DB
    doc_stmt = select(Document).where(Document.id == uuid.UUID(doc_id))
    doc_res = await db_session.execute(doc_stmt)
    doc = doc_res.scalar_one()
    assert doc.status == "chunked"
    assert doc.metadata_ is not None
    assert doc.metadata_["chunk_count"] == summary["total_chunks"]

    # 5. Query chunks API
    list_res = await async_client.get(
        f"/api/v1/documents/{doc_id}/chunks",
        headers=headers,
    )
    assert list_res.status_code == 200
    chunk_data = list_res.json()
    assert chunk_data["total"] == summary["total_chunks"]
    items = chunk_data["items"]
    assert len(items) == summary["total_chunks"]

    # Verify parent and children linkage
    parent_chunks = [c for c in items if c["chunk_type"] == "parent"]
    child_chunks = [c for c in items if c["chunk_type"] != "parent"]
    assert len(parent_chunks) >= 1
    assert len(child_chunks) >= 1

    parent_ids = {p["id"] for p in parent_chunks}
    for child in child_chunks:
        assert child["parent_chunk_id"] in parent_ids
        assert child["token_count"] > 0
        assert child["character_count"] == len(child["content"])
        assert len(child["content_hash"]) == 64
        assert child["chunker_version"] == "1.0.0"


@pytest.mark.asyncio
async def test_idempotent_reprocessing(
    async_client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    """Reprocessing the same document must cleanly replace chunks without duplicates."""
    org = await _create_test_organization(db_session)
    _, token = await _create_test_user_with_role(db_session, org, ROLE_ADMIN)
    headers = {"Authorization": f"Bearer {token}", "X-Organization-ID": str(org.id)}

    docx_bytes = _create_sample_docx()
    upload_res = await async_client.post(
        "/api/v1/documents",
        headers=headers,
        files={
            "file": (
                "reprocess.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    doc_id = upload_res.json()["id"]

    await async_client.post(f"/api/v1/documents/{doc_id}/ingest?sync=true", headers=headers)

    # First chunking run
    res1 = await async_client.post(f"/api/v1/documents/{doc_id}/chunks?sync=true", headers=headers)
    assert res1.status_code == 200
    summary1 = res1.json()

    # Second chunking run (reprocessing)
    res2 = await async_client.post(f"/api/v1/documents/{doc_id}/chunks?sync=true", headers=headers)
    assert res2.status_code == 200
    summary2 = res2.json()

    assert summary1["total_chunks"] == summary2["total_chunks"]

    # Verify DB chunk count directly
    chunks_stmt = select(DocumentChunk).where(DocumentChunk.document_id == uuid.UUID(doc_id))
    chunks_res = await db_session.execute(chunks_stmt)
    persisted_chunks = chunks_res.scalars().all()
    assert len(persisted_chunks) == summary1["total_chunks"]


@pytest.mark.asyncio
async def test_chunk_pagination_and_parent_filtering(
    async_client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    """Test pagination (skip/limit) and parent_only filter on chunks list."""
    org = await _create_test_organization(db_session)
    _, token = await _create_test_user_with_role(db_session, org, ROLE_ADMIN)
    headers = {"Authorization": f"Bearer {token}", "X-Organization-ID": str(org.id)}

    docx_bytes = _create_sample_docx()
    upload_res = await async_client.post(
        "/api/v1/documents",
        headers=headers,
        files={
            "file": (
                "filter.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    doc_id = upload_res.json()["id"]
    await async_client.post(f"/api/v1/documents/{doc_id}/ingest?sync=true", headers=headers)
    await async_client.post(f"/api/v1/documents/{doc_id}/chunks?sync=true", headers=headers)

    # 1. Limit = 2
    paged_res = await async_client.get(
        f"/api/v1/documents/{doc_id}/chunks?skip=0&limit=2",
        headers=headers,
    )
    assert paged_res.status_code == 200
    paged_data = paged_res.json()
    assert len(paged_data["items"]) == 2
    assert paged_data["skip"] == 0
    assert paged_data["limit"] == 2

    # 2. Parent only
    parent_res = await async_client.get(
        f"/api/v1/documents/{doc_id}/chunks?parent_only=true",
        headers=headers,
    )
    assert parent_res.status_code == 200
    parent_data = parent_res.json()
    for item in parent_data["items"]:
        assert item["parent_chunk_id"] is None
        assert item["chunk_type"] == "parent"


@pytest.mark.asyncio
async def test_chunk_quality_summary_endpoint(
    async_client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    """Test GET /api/v1/documents/{id}/chunks/summary returns valid metrics."""
    org = await _create_test_organization(db_session)
    _, token = await _create_test_user_with_role(db_session, org, ROLE_ADMIN)
    headers = {"Authorization": f"Bearer {token}", "X-Organization-ID": str(org.id)}

    docx_bytes = _create_sample_docx()
    upload_res = await async_client.post(
        "/api/v1/documents",
        headers=headers,
        files={
            "file": (
                "summary.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    doc_id = upload_res.json()["id"]
    await async_client.post(f"/api/v1/documents/{doc_id}/ingest?sync=true", headers=headers)
    await async_client.post(f"/api/v1/documents/{doc_id}/chunks?sync=true", headers=headers)

    summary_res = await async_client.get(
        f"/api/v1/documents/{doc_id}/chunks/summary",
        headers=headers,
    )
    assert summary_res.status_code == 200
    summary = summary_res.json()
    assert summary["document_id"] == doc_id
    assert summary["total_chunks"] > 0
    assert summary["mean_tokens"] > 0
    assert "chunker_version" in summary


@pytest.mark.asyncio
async def test_tenant_isolation_cross_tenant_access(
    async_client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    """Org B must not access or chunk documents belonging to Org A."""
    org_a = await _create_test_organization(db_session, "Org A")
    org_b = await _create_test_organization(db_session, "Org B")

    _, token_a = await _create_test_user_with_role(db_session, org_a, ROLE_ADMIN)
    _, token_b = await _create_test_user_with_role(db_session, org_b, ROLE_ADMIN)

    headers_a = {"Authorization": f"Bearer {token_a}", "X-Organization-ID": str(org_a.id)}
    headers_b = {"Authorization": f"Bearer {token_b}", "X-Organization-ID": str(org_b.id)}

    # Org A uploads, ingests, and chunks
    docx_bytes = _create_sample_docx()
    upload_res = await async_client.post(
        "/api/v1/documents",
        headers=headers_a,
        files={
            "file": (
                "isolated.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    doc_id = upload_res.json()["id"]
    await async_client.post(f"/api/v1/documents/{doc_id}/ingest?sync=true", headers=headers_a)
    await async_client.post(f"/api/v1/documents/{doc_id}/chunks?sync=true", headers=headers_a)

    # Org B attempts to trigger chunking on Org A's document
    cross_chunk_res = await async_client.post(
        f"/api/v1/documents/{doc_id}/chunks?sync=true",
        headers=headers_b,
    )
    assert cross_chunk_res.status_code == 404

    # Org B attempts to list chunks of Org A's document
    cross_list_res = await async_client.get(
        f"/api/v1/documents/{doc_id}/chunks",
        headers=headers_b,
    )
    assert cross_list_res.status_code == 404

    # Org B attempts to get summary of Org A's document
    cross_summary_res = await async_client.get(
        f"/api/v1/documents/{doc_id}/chunks/summary",
        headers=headers_b,
    )
    assert cross_summary_res.status_code == 404


@pytest.mark.asyncio
async def test_rbac_permissions(
    async_client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    """Viewer cannot trigger chunking (403), but can read chunks (200). Analyst can chunk (200)."""
    org = await _create_test_organization(db_session)
    _, admin_token = await _create_test_user_with_role(db_session, org, ROLE_ADMIN)
    _, analyst_token = await _create_test_user_with_role(db_session, org, ROLE_ANALYST)
    _, viewer_token = await _create_test_user_with_role(db_session, org, ROLE_VIEWER)

    admin_headers = {"Authorization": f"Bearer {admin_token}", "X-Organization-ID": str(org.id)}
    analyst_headers = {"Authorization": f"Bearer {analyst_token}", "X-Organization-ID": str(org.id)}
    viewer_headers = {"Authorization": f"Bearer {viewer_token}", "X-Organization-ID": str(org.id)}

    # Admin uploads and ingests
    docx_bytes = _create_sample_docx()
    upload_res = await async_client.post(
        "/api/v1/documents",
        headers=admin_headers,
        files={
            "file": (
                "rbac_test.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    doc_id = upload_res.json()["id"]
    await async_client.post(f"/api/v1/documents/{doc_id}/ingest?sync=true", headers=admin_headers)

    # 1. Viewer attempts to trigger chunking -> 403 Forbidden
    viewer_chunk_res = await async_client.post(
        f"/api/v1/documents/{doc_id}/chunks?sync=true",
        headers=viewer_headers,
    )
    assert viewer_chunk_res.status_code == 403

    # 2. Analyst triggers chunking -> 200 OK
    analyst_chunk_res = await async_client.post(
        f"/api/v1/documents/{doc_id}/chunks?sync=true",
        headers=analyst_headers,
    )
    assert analyst_chunk_res.status_code == 200

    # 3. Viewer reads chunks -> 200 OK
    viewer_read_res = await async_client.get(
        f"/api/v1/documents/{doc_id}/chunks",
        headers=viewer_headers,
    )
    assert viewer_read_res.status_code == 200


@pytest.mark.asyncio
async def test_chunking_invalid_status_failure(
    async_client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    """Cannot chunk a document that is in 'uploaded' status without being parsed."""
    org = await _create_test_organization(db_session)
    _, token = await _create_test_user_with_role(db_session, org, ROLE_ADMIN)
    headers = {"Authorization": f"Bearer {token}", "X-Organization-ID": str(org.id)}

    docx_bytes = _create_sample_docx()
    upload_res = await async_client.post(
        "/api/v1/documents",
        headers=headers,
        files={
            "file": (
                "unparsed.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    doc_id = upload_res.json()["id"]

    # Ensure document status is 'uploaded' and not auto-ingested
    doc_stmt = select(Document).where(Document.id == uuid.UUID(doc_id))
    doc_res = await db_session.execute(doc_stmt)
    doc = doc_res.scalar_one()
    doc.status = "uploaded"
    await db_session.commit()

    # Attempt to chunk directly before ingestion
    chunk_res = await async_client.post(
        f"/api/v1/documents/{doc_id}/chunks?sync=true",
        headers=headers,
    )
    assert chunk_res.status_code == 400
    assert chunk_res.json()["error"]["code"] == "INVALID_DOCUMENT_STATUS"


@pytest.mark.asyncio
async def test_cascading_chunk_deletion(
    async_client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    """Deleting a document should remove or soft-delete all linked chunks."""
    org = await _create_test_organization(db_session)
    _, token = await _create_test_user_with_role(db_session, org, ROLE_ADMIN)
    headers = {"Authorization": f"Bearer {token}", "X-Organization-ID": str(org.id)}

    docx_bytes = _create_sample_docx()
    upload_res = await async_client.post(
        "/api/v1/documents",
        headers=headers,
        files={
            "file": (
                "cascade.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    doc_id = upload_res.json()["id"]
    await async_client.post(f"/api/v1/documents/{doc_id}/ingest?sync=true", headers=headers)
    chunk_res = await async_client.post(
        f"/api/v1/documents/{doc_id}/chunks?sync=true", headers=headers
    )
    assert chunk_res.status_code == 200

    # Verify chunks exist
    chunks_stmt = select(DocumentChunk).where(DocumentChunk.document_id == uuid.UUID(doc_id))
    chunks_before = (await db_session.execute(chunks_stmt)).scalars().all()
    assert len(chunks_before) > 0

    # Delete document
    del_res = await async_client.delete(f"/api/v1/documents/{doc_id}", headers=headers)
    assert del_res.status_code == 200

    # Verify querying chunks via API returns 404
    get_res = await async_client.get(f"/api/v1/documents/{doc_id}/chunks", headers=headers)
    assert get_res.status_code == 404
