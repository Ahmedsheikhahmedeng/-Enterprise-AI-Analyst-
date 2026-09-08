"""Integration tests for Data Connectors, Schema Discovery, Agent Tools, and Jobs."""

import uuid
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import openpyxl
import pytest
from sqlalchemy.ext.asyncio import AsyncSession

from app.agents.context import AgentExecutionContext
from app.agents.tools.datasource import (
    DataSourceListTool,
    DataSourceQueryTool,
)
from app.connectors.application.schema_service import SchemaService
from app.connectors.domain.enums import ConnectorType, SyncStatus, SyncType
from app.connectors.domain.models import ColumnSchema, SchemaModel, TableSchema
from app.connectors.infrastructure.csv.connector import CSVConnector
from app.connectors.infrastructure.excel.connector import ExcelConnector
from app.jobs.context import JobContext
from app.jobs.tasks.sync import DataSourceSyncTask
from app.models.data_source import DataSource


@pytest.mark.asyncio
async def test_csv_integration_full_flow(tmp_path: Path) -> None:
    csv_file = tmp_path / "customers.csv"
    csv_file.write_text(
        "customer_id,name,tier,balance\n1,Acme Corp,Enterprise,95000.50\n2,Globex,Growth,1250.00\n"
    )

    conn = CSVConnector()
    ds_id = uuid.uuid4()
    org_id = uuid.uuid4()
    config = {"file_path": str(csv_file)}

    # 1. Connection test
    test_res = await conn.test_connection(config)
    assert test_res.success is True
    assert test_res.latency_ms >= 0

    # 2. Schema discovery
    schema = await conn.get_schema(ds_id, org_id, config)
    assert len(schema.tables) == 1
    assert schema.tables[0].name == "customers"
    assert len(schema.tables[0].columns) == 4

    # 3. Preview
    preview = await conn.preview_data(ds_id, org_id, config, max_rows=10)
    assert len(preview.rows) == 2
    assert preview.rows[0]["name"] == "Acme Corp"


@pytest.mark.asyncio
async def test_excel_integration_full_flow(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Q1_Financials"
    ws1.append(["metric", "value", "status"])
    ws1.append(["ARR", 5000000, "on_track"])
    ws1.append(["Churn", 0.02, "healthy"])

    ws2 = wb.create_sheet(title="Headcount")
    ws2.append(["department", "count"])
    ws2.append(["Engineering", 42])

    excel_file = tmp_path / "financials.xlsx"
    wb.save(str(excel_file))

    conn = ExcelConnector()
    ds_id = uuid.uuid4()
    org_id = uuid.uuid4()
    config = {"file_path": str(excel_file)}

    # 1. Connection test
    test_res = await conn.test_connection(config)
    assert test_res.success is True

    # 2. Schema discovery across sheets
    schema = await conn.get_schema(ds_id, org_id, config)
    sheet_names = [t.name for t in schema.tables]
    assert "Q1_Financials" in sheet_names
    assert "Headcount" in sheet_names

    # 3. Preview specific sheet
    preview = await conn.preview_data(
        ds_id, org_id, config, target_name="Q1_Financials", max_rows=10
    )
    assert len(preview.rows) == 2
    assert preview.rows[0]["metric"] == "ARR"


@pytest.mark.asyncio
async def test_schema_service_redis_caching_and_invalidation() -> None:
    org_id = uuid.uuid4()
    ds_id = uuid.uuid4()

    mock_redis = AsyncMock()
    cached_data: dict[str, str] = {}

    async def mock_get(key: str) -> str | None:
        return cached_data.get(key)

    async def mock_set(key: str, val: str, ex: int | None = None) -> None:
        cached_data[key] = val

    async def mock_keys(pattern: str) -> list[str]:
        prefix = pattern.replace("*", "")
        return [k for k in cached_data if k.startswith(prefix)]

    async def mock_delete(*keys: str) -> None:
        for k in keys:
            cached_data.pop(k, None)

    mock_redis.get.side_effect = mock_get
    mock_redis.set.side_effect = mock_set
    mock_redis.keys.side_effect = mock_keys
    mock_redis.delete.side_effect = mock_delete

    service = SchemaService(redis_client=mock_redis)

    col = ColumnSchema(name="id", data_type="integer", nullable=False, primary_key=True)
    table = TableSchema(name="items", columns=[col])
    schema = SchemaModel(
        datasource_id=ds_id,
        organization_id=org_id,
        connector_type=ConnectorType.POSTGRESQL,
        tables=[table],
    )

    # 1. Cache schema
    await service.cache_schema(schema, version=1)
    key = service._cache_key(org_id, ds_id, version=1)
    assert key in cached_data

    # 2. Read cached schema
    retrieved = await service.get_cached_schema(org_id, ds_id, version=1)
    assert retrieved is not None
    assert retrieved.tables[0].name == "items"

    # 3. Invalidate cache
    await service.invalidate_schema_cache(org_id, ds_id)
    assert key not in cached_data


@pytest.mark.asyncio
async def test_agent_runtime_datasource_tools() -> None:
    org_id = uuid.uuid4()
    ds_id = uuid.uuid4()
    user_id = uuid.uuid4()

    mock_session = AsyncMock(spec=AsyncSession)
    mock_ds = DataSource(
        id=ds_id,
        organization_id=org_id,
        name="Sales DB",
        type=ConnectorType.POSTGRESQL.value,
        status="active",
        configuration={"host": "localhost", "port": 5432, "dbname": "test"},
    )
    mock_result = MagicMock()
    mock_result.scalars.return_value.all.return_value = [mock_ds]
    mock_result.scalars.return_value.first.return_value = mock_ds
    mock_session.execute.return_value = mock_result

    context = AgentExecutionContext(
        organization_id=org_id,
        session_id=uuid.uuid4(),
        step_id=uuid.uuid4(),
        db_session=mock_session,
        user_id=user_id,
    )

    # 1. List tool
    list_tool = DataSourceListTool()
    validated_list = await list_tool.validate({"limit": 10, "offset": 0})
    res_list = await list_tool.execute(validated_list, context)
    assert res_list.success is True

    # 2. Query tool with safe SQL
    query_tool = DataSourceQueryTool()
    mock_connector = MagicMock()
    from app.connectors.domain.capabilities import ConnectorCapabilities

    mock_connector.capabilities = ConnectorCapabilities(query=True, schema_read=True)
    from app.connectors.domain.models import QueryResult

    mock_connector.execute_query = AsyncMock(
        return_value=QueryResult(
            columns=["id", "total"],
            rows=[{"id": 1, "total": 100}],
            row_count=1,
            execution_time_ms=10.0,
            datasource_id=ds_id,
            provenance={},
        )
    )
    query_tool.query_service.registry.register(
        ConnectorType.POSTGRESQL.value, lambda: mock_connector
    )

    validated_query = await query_tool.validate(
        {
            "datasource_id": str(ds_id),
            "query": "SELECT id, total FROM orders LIMIT 10",
        }
    )
    res_query = await query_tool.execute(validated_query, context)
    assert res_query.success is True
    assert res_query.data["row_count"] == 1


@pytest.mark.asyncio
async def test_background_job_sync_task() -> None:
    org_id = uuid.uuid4()
    ds_id = uuid.uuid4()
    job_id = uuid.uuid4()

    mock_sync_service = AsyncMock()
    from app.connectors.domain.models import SyncRunModel

    mock_sync_service.execute_sync.return_value = SyncRunModel(
        id=uuid.uuid4(),
        data_source_id=ds_id,
        organization_id=org_id,
        sync_type=SyncType.FULL,
        status=SyncStatus.COMPLETED,
        started_at=datetime.now(UTC),
        completed_at=datetime.now(UTC),
        rows_synced=150,
    )

    mock_session = AsyncMock(spec=AsyncSession)
    from collections.abc import AsyncIterator
    from contextlib import asynccontextmanager

    @asynccontextmanager
    async def mock_session_factory() -> AsyncIterator[AsyncSession]:
        yield mock_session

    task = DataSourceSyncTask(
        sync_service=mock_sync_service,
        session_factory=mock_session_factory,  # type: ignore[arg-type]
    )
    context = JobContext(
        job_id=job_id,
        organization_id=org_id,
        job_type="data_source_sync",
        attempt=1,
        max_attempts=3,
    )

    output = await task.run(
        payload={"datasource_id": str(ds_id), "sync_type": "full"},
        context=context,
    )

    assert output["status"].lower() == "completed"
    assert output["records_synced"] == 150
