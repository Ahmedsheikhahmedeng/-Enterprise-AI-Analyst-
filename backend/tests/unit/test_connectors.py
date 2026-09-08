"""Unit tests for Connector domain models, SecretProvider, Registry, and Connectors."""

import uuid
from pathlib import Path

import openpyxl
import pytest

from app.connectors.domain.capabilities import ConnectorCapabilities
from app.connectors.domain.enums import ConnectionStatus, ConnectorType
from app.connectors.domain.errors import (
    InvalidConfigurationError,
    UnsupportedCapabilityError,
)
from app.connectors.domain.models import (
    ColumnSchema,
    QueryRequest,
    QueryResult,
    SchemaModel,
    TableSchema,
)
from app.connectors.infrastructure.csv.connector import CSVConnector
from app.connectors.infrastructure.excel.connector import ExcelConnector
from app.connectors.infrastructure.postgres.connector import PostgreSQLConnector
from app.connectors.infrastructure.registry import ConnectorRegistry
from app.connectors.infrastructure.secrets import SecretProvider


def test_connector_capabilities_enforcement() -> None:
    caps = ConnectorCapabilities(
        schema_read=True,
        query=False,
        write=False,
        sync=False,
    )
    caps.require_schema()  # Should not raise

    with pytest.raises(UnsupportedCapabilityError) as exc:
        caps.require_query()
    assert "query" in str(exc.value)

    with pytest.raises(UnsupportedCapabilityError) as exc:
        caps.require_sync()
    assert "sync" in str(exc.value)


def test_connector_registry() -> None:
    reg = ConnectorRegistry()
    assert reg.has("postgresql")
    assert reg.has("csv")
    assert reg.has("excel")
    assert not reg.has("unsupported_db")

    pg = reg.get("postgresql")
    assert isinstance(pg, PostgreSQLConnector)
    assert pg.capabilities.query is True

    csv_conn = reg.get("csv")
    assert isinstance(csv_conn, CSVConnector)
    assert csv_conn.capabilities.files is True

    excel_conn = reg.get("excel")
    assert isinstance(excel_conn, ExcelConnector)
    assert excel_conn.capabilities.files is True

    with pytest.raises(InvalidConfigurationError):
        reg.get("unknown_type")


def test_secret_provider_encryption_and_redaction() -> None:
    provider = SecretProvider(secret_key="my-super-secret-key-for-unit-testing")

    config = {
        "host": "localhost",
        "port": 5432,
        "dbname": "analytics",
        "user": "db_user",
        "password": "SuperSecretPassword123!",
        "api_key": "mock_api_key_1234567890",
        "token": "tok_abcdef",
    }

    encrypted = provider.encrypt_config(config)
    # Ensure sensitive fields are encrypted and not plaintext
    assert encrypted["password"] != "SuperSecretPassword123!"
    assert encrypted["api_key"] != "mock_api_key_1234567890"
    assert encrypted["token"] != "tok_abcdef"
    assert encrypted["host"] == "localhost"

    # Decrypt and verify fidelity
    decrypted = provider.decrypt_config(encrypted)
    assert decrypted["password"] == "SuperSecretPassword123!"
    assert decrypted["api_key"] == "mock_api_key_1234567890"
    assert decrypted["token"] == "tok_abcdef"

    # Redact and verify zero secret leakage
    redacted = provider.redact_config(decrypted)
    assert redacted["password"] == "***REDACTED***"
    assert redacted["api_key"] == "***REDACTED***"
    assert redacted["token"] == "***REDACTED***"
    assert redacted["host"] == "localhost"
    assert redacted["user"] == "db_user"


def test_schema_model_serialization() -> None:
    org_id = uuid.uuid4()
    ds_id = uuid.uuid4()
    col1 = ColumnSchema(name="id", data_type="integer", nullable=False, primary_key=True)
    col2 = ColumnSchema(name="email", data_type="varchar", nullable=True)
    table = TableSchema(name="users", columns=[col1, col2], primary_keys=["id"])

    schema = SchemaModel(
        datasource_id=ds_id,
        organization_id=org_id,
        connector_type=ConnectorType.POSTGRESQL,
        tables=[table],
    )

    data = schema.model_dump()
    assert data["datasource_id"] == ds_id
    assert data["tables"][0]["name"] == "users"
    assert len(data["tables"][0]["columns"]) == 2

    # Round trip
    parsed = SchemaModel.model_validate(data)
    assert parsed.tables[0].columns[0].primary_key is True


def test_query_request_and_result() -> None:
    ds_id = uuid.uuid4()
    org_id = uuid.uuid4()
    req = QueryRequest(
        datasource_id=ds_id,
        organization_id=org_id,
        query="SELECT id, name FROM customers LIMIT 10",
        timeout_ms=5000,
        max_rows=100,
    )
    assert req.datasource_id == ds_id

    res = QueryResult(
        columns=["id", "name"],
        rows=[{"id": 1, "name": "Alice"}, {"id": 2, "name": "Bob"}],
        row_count=2,
        execution_time_ms=12.5,
        datasource_id=ds_id,
        provenance={"query_hash": "abc123hash"},
    )
    assert res.row_count == 2
    assert res.columns == ["id", "name"]


@pytest.mark.asyncio
async def test_csv_connector_schema_and_preview(tmp_path: Path) -> None:
    csv_file = tmp_path / "test_data.csv"
    csv_file.write_text("id,name,amount,active\n1,Alice,100.5,true\n2,Bob,250.0,false\n")

    conn = CSVConnector()
    ds_id = uuid.uuid4()
    org_id = uuid.uuid4()
    config = {"file_path": str(csv_file)}

    test_res = await conn.test_connection(config)
    assert test_res.success is True

    schema = await conn.get_schema(ds_id, org_id, config)
    assert len(schema.tables) == 1
    assert schema.tables[0].name == "test_data"
    col_names = [c.name for c in schema.tables[0].columns]
    assert col_names == ["id", "name", "amount", "active"]

    preview = await conn.preview_data(ds_id, org_id, config, max_rows=10)
    assert len(preview.rows) == 2
    preview_col_names = [c.name for c in preview.columns]
    assert preview_col_names == ["id", "name", "amount", "active"]
    assert preview.rows[0]["name"] == "Alice"


@pytest.mark.asyncio
async def test_csv_formula_injection_defense(tmp_path: Path) -> None:
    csv_file = tmp_path / "malicious.csv"
    csv_file.write_text("id,cmd,payload\n1,=cmd|' /C calc'!A0,@SUM(1+1)\n2,-2+5,+cmd\n")

    conn = CSVConnector()
    ds_id = uuid.uuid4()
    org_id = uuid.uuid4()
    config = {"file_path": str(csv_file)}

    preview = await conn.preview_data(ds_id, org_id, config, max_rows=5)
    # Malicious leading formula characters should be neutralized with single quote prefix
    assert preview.rows[0]["cmd"].startswith("'=")
    assert preview.rows[0]["payload"].startswith("'@")
    assert preview.rows[1]["cmd"].startswith("'-")
    assert preview.rows[1]["payload"].startswith("'+")


@pytest.mark.asyncio
async def test_excel_connector_schema_and_preview(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sales"
    ws.append(["order_id", "product", "revenue"])
    ws.append([101, "Laptop", 1200.50])
    ws.append([102, "Mouse", 25.00])
    excel_path = tmp_path / "test_workbook.xlsx"
    wb.save(str(excel_path))

    conn = ExcelConnector()
    ds_id = uuid.uuid4()
    org_id = uuid.uuid4()
    config = {"file_path": str(excel_path)}

    test_res = await conn.test_connection(config)
    assert test_res.success is True

    schema = await conn.get_schema(ds_id, org_id, config)
    assert len(schema.tables) == 1
    assert schema.tables[0].name == "Sales"
    col_names = [c.name for c in schema.tables[0].columns]
    assert col_names == ["order_id", "product", "revenue"]

    preview = await conn.preview_data(ds_id, org_id, config, target_name="Sales", max_rows=5)
    assert len(preview.rows) == 2
    assert preview.rows[0]["product"] == "Laptop"


@pytest.mark.asyncio
async def test_excel_formula_injection_defense(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Formulas"
    ws.append(["id", "formula_field"])
    ws.append([1, "=1+1"])
    ws.append([2, "@calc"])
    excel_path = tmp_path / "formula.xlsx"
    wb.save(str(excel_path))

    conn = ExcelConnector()
    ds_id = uuid.uuid4()
    org_id = uuid.uuid4()
    config = {"file_path": str(excel_path)}

    preview = await conn.preview_data(ds_id, org_id, config, max_rows=5)
    # Formula characters must be prepended with quote
    assert preview.rows[0]["formula_field"].startswith("'=")
    assert preview.rows[1]["formula_field"].startswith("'@")


def test_connection_status_enum_values() -> None:
    assert ConnectionStatus.REGISTERED.value == "registered"
    assert ConnectionStatus.CONNECTING.value == "connecting"
    assert ConnectionStatus.ACTIVE.value == "active"
    assert ConnectionStatus.DEGRADED.value == "degraded"
    assert ConnectionStatus.ERROR.value == "error"
    assert ConnectionStatus.DISABLED.value == "disabled"
